import json, pathlib, sys
from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium
from openpyxl import load_workbook
from pptx import Presentation

root = pathlib.Path(sys.argv[1]).resolve()
checks = json.loads(pathlib.Path(sys.argv[2]).read_text(encoding='utf-8'))
result = []
for item in checks:
    file = (root / item['path']).resolve()
    assert file.is_relative_to(root)
    values = []
    if file.suffix == '.docx':
        doc = Document(file)
        values.extend(p.text for p in doc.paragraphs)
        values.extend(c.text for table in doc.tables for row in table.rows for c in row.cells)
    elif file.suffix == '.pdf':
        reader = PdfReader(file)
        assert all(abs(float(page.mediabox.width)-595.28)<2 and abs(float(page.mediabox.height)-841.89)<2 for page in reader.pages)
        # pypdf inserts inferred spaces between bold Chromium glyphs. PDFium
        # resolves their actual positions; assert the exact text without stripping.
        document = pdfium.PdfDocument(file)
        try:
            for index in range(len(document)):
                page = document[index]
                text_page = page.get_textpage()
                try:
                    values.append(text_page.get_text_range())
                finally:
                    text_page.close()
                    page.close()
        finally:
            document.close()
    elif file.suffix == '.xlsx':
        book = load_workbook(file, data_only=False)
        values.extend(cell.value for sheet in book for row in sheet for cell in row if cell.value is not None)
        assert not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row)
        if item.get('number') is not None:
            assert item['number'] in values
    elif file.suffix == '.pptx':
        deck = Presentation(file)
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    values.append(shape.text)
                if shape.has_table:
                    values.extend(cell.text for row in shape.table.rows for cell in row.cells)
    else:
        raise AssertionError('unsupported output')
    text = '\n'.join(str(v) for v in values)
    for marker in item['markers']:
        assert marker in text, f'{file.name}: expected marker missing'
    result.append({'path': file.name, 'format': file.suffix[1:], 'markersPresent': True, 'bytes': file.stat().st_size})
print(json.dumps({'success': True, 'files': result}, ensure_ascii=False))
