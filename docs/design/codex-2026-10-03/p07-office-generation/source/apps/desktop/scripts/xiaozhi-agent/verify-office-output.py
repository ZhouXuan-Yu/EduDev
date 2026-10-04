import json,sys,pathlib
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
root=pathlib.Path(sys.argv[1]);marker=sys.argv[2]
doc=Document(root/'generated.docx')
doc_text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
assert marker in doc_text and '最后一段中文验收' in doc_text and '=1+1' in doc_text
assert len(doc.tables)==1 and doc.sections[0].page_width.twips==11906 and doc.sections[0].page_height.twips==16838
wb=load_workbook(root/'generated.xlsx',data_only=False)
cells=[c for sheet in wb for row in sheet for c in row if c.value is not None]
assert any(c.value==37 and c.data_type=='n' for c in cells) and any(c.value=='=1+1' and c.data_type=='s' for c in cells)
assert not any(c.data_type=='f' for c in cells)
xtext='\n'.join(str(c.value) for c in cells);assert marker in xtext and '最后一段中文验收' in xtext
assert len(wb.worksheets)==2 and all(s.page_setup.paperSize==9 for s in wb)
ppt=Presentation(root/'generated.pptx')
texts=[]
for slide in ppt.slides:
 for shape in slide.shapes:
  if shape.has_text_frame:texts.append(shape.text)
  if shape.has_table:texts.extend(c.text for r in shape.table.rows for c in r.cells)
ptext='\n'.join(texts)
assert marker in ptext and '最后一段中文验收' in ptext and '=1+1' in ptext
assert '数据末行100' in ptext and len(ppt.slides)>20
pdf=PdfReader(root/'generated.pdf')
pdftext='\n'.join(p.extract_text() for p in pdf.pages)
assert marker in pdftext and '最后一段中文验收' in pdftext and '分数' in pdftext and '数据末行100' in pdftext
assert len(pdf.pages)>1
for page in pdf.pages:
 assert abs(float(page.mediabox.width)-595.28)<1 and abs(float(page.mediabox.height)-841.89)<1
assert '<script>' in pdftext and '127.0.0.1' in pdftext
result={'success':True,'docx':{'tables':len(doc.tables),'a4':True},'xlsx':{'sheets':len(wb.worksheets),'numeric37':True,'formulaIsLiteral':True},'pptx':{'slides':len(ppt.slides),'lastTableRow':True},'pdf':{'pages':len(pdf.pages),'a4':True,'chineseAndTail':True},'boundary':'Actual generated formats independently parsed; not real teacher UI or WPS/Office rendering proof'}
(root/'format-readback.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(result,ensure_ascii=False))
