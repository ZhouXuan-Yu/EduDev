"""Synthetic acceptance input only; no product Python runtime dependency."""
import sys, json, hashlib
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont

root=Path(sys.argv[1]).resolve()
root.mkdir(parents=True, exist_ok=True)
marker=sys.argv[2]
doc=Document()
doc.add_heading('分数教研记录',0)
doc.add_paragraph(f'{marker} 课堂37分钟，教师可修改。')
table=doc.add_table(rows=1,cols=2)
table.rows[0].cells[0].text='知识点';table.rows[0].cells[1].text='练习数量'
row=table.add_row().cells;row[0].text='分数';row[1].text='8'
doc.add_paragraph('https://example.invalid/no-fetch')
doc.save(root/'备课.docx')
book=Workbook();sheet=book.active;sheet.title='课时'
sheet.append(['标识','课时','练习数量']);sheet.append([marker,37,8])
sheet['D1']='公式';sheet['D2']='=B2+C2'
book.save(root/'教研.xlsx')
slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text='教学课件'
slide.placeholders[1].text=f'{marker}\n分数课堂37分钟\n练习8题'
slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='第二课时';slide.placeholders[1].text='教师确认后实施'
slides.save(root/'课件.pptx')
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
cid=canvas.Canvas(str(root/'缺文字映射.pdf'),pagesize=A4)
cid.setFont('STSong-Light',14);cid.drawString(50,790,'分数课堂讲义');cid.save()
pdfmetrics.registerFont(TTFont('FixtureChinese','C:/Windows/Fonts/simsun.ttc',subfontIndex=0))
pdf=canvas.Canvas(str(root/'讲义.pdf'),pagesize=A4)
pdf.setFont('FixtureChinese',14);pdf.drawString(50,790,'分数课堂讲义')
pdf.setFont('Helvetica',12);pdf.drawString(50,750,f'{marker} Lesson 37 minutes, exercises 8')
pdf.showPage();pdf.drawString(50,750,'Second page: teacher review');pdf.save()
scan=canvas.Canvas(str(root/'扫描.pdf'),pagesize=A4)
scan.drawImage(str(Path(__file__).resolve().parents[4]/'docs/design/codex-2026-10-03/process-reference.png'),50,120,width=400,height=600)
scan.save()
(root/'损坏.docx').write_bytes(b'not a document')
(root/'备注.txt').write_text(marker+' local text',encoding='utf8')
(root/'.env.local').write_text('SYNTHETIC_SECRET',encoding='utf8')
names=['备课.docx','教研.xlsx','课件.pptx','讲义.pdf','扫描.pdf','缺文字映射.pdf','损坏.docx','备注.txt']
facts={name:hashlib.sha256((root/name).read_bytes()).hexdigest() for name in names}
(root/'fixture-facts.json').write_text(json.dumps({'marker':marker,'hashes':facts},ensure_ascii=False),encoding='utf8')
print('Generated four real Office/PDF fixtures plus scanned and malformed cases')
