"""Read the actual WPS COM result against owned fixture values, independently of PDF markers."""
import json,sys,pathlib,collections,hashlib
from docx import Document
from openpyxl import load_workbook

root=pathlib.Path(sys.argv[1]);source=pathlib.Path(sys.argv[2])
draft=json.loads((source/('fixture.json' if (source/'fixture.json').exists() else 'draft.json')).read_text(encoding='utf-8'))
report=json.loads((root/'report.json').read_text(encoding='utf-8-sig'))
assert report['success'] and all(a['owned'] and a['inputUnchanged'] for a in report['applications'])
checks=[]
native={a['kind']:a for a in report['applications']}
document=Document(source/'generated.docx')
paragraphs=[draft['title']]
for section in draft['sections']:paragraphs+=[section['heading'],*section['paragraphs']]
assert [p.text for p in document.paragraphs]==paragraphs
native_word=native['docx']['text'].replace('\v','\n') # Actual Word manual line break, not a paragraph boundary.
for text in paragraphs:assert text in native_word
tables=[s['table'] for s in draft['sections'] if 'table' in s]
assert len(document.tables)==len(tables)==native['docx']['tables']
for actual,table in zip(document.tables,tables):
 expected=[table['columns'],*[[str(v) for v in row] for row in table['rows']]]
 assert [[c.text for c in row.cells] for row in actual.rows]==expected
 for row in expected:
  for cell in row:assert cell in native_word
checks.append('Actual WPS Word content and independent DOCX all paragraphs/table cells preserve exact spaces and values')
workbook=load_workbook(source/'generated.xlsx',data_only=False)
assert len(workbook.worksheets)==len(native['xlsx']['values'])
for sheet,actual in zip(workbook,native['xlsx']['values']):
 expected=[c.value for row in sheet for c in row]
 assert expected==actual
 for row in sheet:
  for cell in row:
   assert cell.data_type!='f'
   if cell.value is not None:assert type(cell.value)==type(actual[(cell.row-1)*sheet.max_column+cell.column-1]) or isinstance(cell.value,(int,float)) and isinstance(actual[(cell.row-1)*sheet.max_column+cell.column-1],(int,float))
checks.append('Actual WPS Excel complete used range equals OOXML, preserving numeric types, literal digits and no formulas')
clean=lambda text:''.join(text.splitlines())
paragraph_cells=[];table_cells=[];columns=tables[0]['columns'] if tables else []
for table in native['pptx']['tables']:
 rows=table['rows']
 if table['columns']==1:paragraph_cells.extend(clean(row[0]) for row in rows)
 else:
  assert columns and [clean(c) for c in rows[0]]==columns
  table_cells.extend(clean(c) for row in rows[1:] for c in row)
expected_paragraphs=[clean(p) for s in draft['sections'] for p in s['paragraphs']]
assert ''.join(paragraph_cells)==''.join(expected_paragraphs)
# SDK may continue one row on a later slide. Compare per-character multiplicity
# across actual native cells, retaining literal spaces; the visual pass checks order/layout.
assert collections.Counter(''.join(table_cells))==collections.Counter(''.join(clean(str(c)) for t in tables for row in t['rows'] for c in row))
for text in [draft['title'],*[s['heading'] for s in draft['sections']]]:assert clean(text) in [clean(t) for t in native['pptx']['text']]
checks.append('Actual WPS presentation all paragraph characters/ordered text and table character counts retained across original autoPage fragments')
result={'success':True,'checks':checks,'sourceHashes':{k:hashlib.sha256((source/('generated.'+k)).read_bytes()).hexdigest() for k in ['docx','xlsx','pptx']},'boundary':'Exact Word/Excel content from actual WPS COM; PPT paragraphs ordered, table fragments all character counts incl. spaces, visual inspection still required for ordering/clipping.'}
(root/'content-readback.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False))
