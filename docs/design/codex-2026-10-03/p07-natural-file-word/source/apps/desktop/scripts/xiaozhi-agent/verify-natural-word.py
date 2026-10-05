"""Independent readback of approved synthetic Word; optional real WPS evidence."""
import json
import sys
from pathlib import Path
from docx import Document

source, expected = Path(sys.argv[1]), json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
document = Document(source)
actual_paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
required = [expected["title"]]
tables = []
for section in expected["sections"]:
    required += [section["heading"], *section["paragraphs"]]
    if "table" in section:
        table = section["table"]
        tables.append([table["columns"], *[[str(v) for v in row] for row in table["rows"]]])
assert actual_paragraphs == required, (actual_paragraphs, required)
actual_tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
assert actual_tables == tables, (actual_tables, tables)
result = {"success": True, "paragraphs": len(actual_paragraphs), "tables": len(actual_tables), "rows": [len(t) for t in actual_tables]}
if len(sys.argv) > 3:
    wps = json.loads(Path(sys.argv[3]).read_text(encoding="utf-8-sig"))
    assert wps["success"] and len(wps["applications"]) == 1
    native = wps["applications"][0]
    assert native["kind"] == "docx" and native["owned"] and native["inputUnchanged"]
    assert native["tables"] == len(tables) and native["pages"] > 0 and native["renderBytes"] > 0
    flattened = required + [cell for table in tables for row in table for cell in row if cell.strip()]
    for text in flattened:
        assert text in native["text"], text
    edited = Document(native["editCopy"]["path"])
    marker = native["editCopy"]["marker"]
    assert marker in edited.tables[0].rows[1].cells[-1].text
    assert marker in native["editCopy"]["reopenedText"]
    result["wps"] = {"version": native["version"], "pages": native["pages"], "editSavedAndReopened": True, "originalUnchanged": True}
print(json.dumps(result, ensure_ascii=False))
