"""Independent readback of actual approved v4 Office outputs; not a correctness grader."""
import io,json,pathlib,sys,zipfile,re,hashlib,xml.etree.ElementTree as ET
from PIL import Image,ImageStat
from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium
report=json.loads(pathlib.Path(sys.argv[1]).read_text(encoding='utf-8'));out=pathlib.Path(sys.argv[1]).parent;root=pathlib.Path(report['result']['root']);checks=[]
ns={'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
def plain_parts(text):
    return re.split(r'\$\$[\s\S]*?\$\$|\$[^$]*?\$|\\\([\s\S]*?\\\)|\\\[[\s\S]*?\\\]',text)
def compact(v):return ''.join(v.split())
def pixels(data):
    with Image.open(io.BytesIO(data)) as image:
        rgb=image.convert('RGB');return rgb.size,hashlib.sha256(rgb.tobytes()).hexdigest()
for item in report['result']['deliveries']:
    file=pathlib.Path(item.get('root',root))/item['path'];draft=item['draft'];paragraphs=[p for s in draft['sections'] for p in s['paragraphs']];math=[p for p in paragraphs if '$' in p or '\\(' in p or '\\[' in p];rendered=[]
    figures=[f for s in draft['sections'] for f in s.get('figures',[])];source_images=[]
    for figure in figures:
        source=(out/'data/question-assets/v1'/f"{figure['assetId']}.png").read_bytes();assert hashlib.sha256(source).hexdigest()==figure['sha256'];source_images.append(source)
    if item['format']=='docx':
        document=Document(file);text='\n'.join(p.text for p in document.paragraphs)
        with zipfile.ZipFile(file) as z:
            xml=ET.fromstring(z.read('word/document.xml'));descriptions=[e.attrib['descr'] for e in xml.findall('.//wp:docPr',ns) if e.attrib.get('name','').startswith('math-')];assert [compact(v) for v in descriptions]==[compact(v) for v in math],'Missing/reordered original formula alt text'
            assert len(document.inline_shapes)==len(math)+len(figures)
            embedded=[z.read(n) for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')]
            for source in source_images:assert source in embedded,'Word source PNG changed or missing'
            for n,name in enumerate([n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')]):
                data=z.read(name)
                if data in source_images:continue
                im=Image.open(io.BytesIO(data)).convert('RGB');assert im.width==1240 and im.height>10 and min(ImageStat.Stat(im).stddev)>2,'Blank/wrong raster'
                target=out/(file.stem+'-formula-'+str(n+1)+'.png');im.save(target);rendered.append(str(target))
        assert all(abs(s.page_width.mm-210)<1 and abs(s.page_height.mm-297)<1 for s in document.sections)
        for p in paragraphs:
            assert compact(p) in [compact(v) for v in descriptions] or compact(p) in compact(text),'Unmatched full paragraph'
        text+='\n'+'\n'.join(descriptions)
    else:
        reader=PdfReader(file);text='\n'.join(p.extract_text() for p in reader.pages);doc=pdfium.PdfDocument(file)
        embedded=[image.data for page in reader.pages for image in page.images]
        for source in source_images:assert pixels(source) in [pixels(data) for data in embedded],'PDF source image pixels changed or missing'
        for i in range(len(doc)):
            page=doc[i];bitmap=page.render(scale=1.5);target=out/(file.stem+'-page-'+str(i+1)+'.png');bitmap.to_pil().save(target);bitmap.close();page.close();rendered.append(str(target))
        doc.close()
        assert all(abs(float(p.mediabox.width)-595.28)<2 and abs(float(p.mediabox.height)-841.89)<2 for p in reader.pages)
        for p in paragraphs:
            for fragment in plain_parts(p):
                if len(compact(fragment))>4:assert compact(fragment) in compact(text),'Missing plain part: '+fragment
        assert not any('$$' in (p.extract_text() or '') for p in reader.pages),'Raw formula delimiter leaked'
    assert '满分100分' in compact(text)
    if item['audience']=='student':assert '参考答案' not in text and '解析尚未独立核验' not in text
    else:assert '尚未独立核验' in text
    checks.append({'path':item['path'],'audience':item['audience'],'mathParagraphs':len(math),'sourceImages':len(figures),'bytes':file.stat().st_size,'a4':True,'rendered':rendered})
(out/'formula-files-readback.json').write_text(json.dumps({'success':True,'files':checks},ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({'success':True,'files':len(checks)},ensure_ascii=False))

