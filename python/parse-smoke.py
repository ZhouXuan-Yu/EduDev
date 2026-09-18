"""Deterministic smoke for the local document parsing worker.

Builds a real PDF (hand-written xref) and a Chinese DOCX fixture, then runs the
worker protocol in-process and asserts extraction plus error paths.
Run: python parse-smoke.py <output-dir>
"""
import json
import sys
import tempfile
from pathlib import Path

from omni_edu_parse.parse import _handle

PDF_TEXT = "Slope of a linear function y = 2x + 3"
DOCX_TEXT = "一次函数斜率教学讲义"
DOCX_BODY = "斜率的定义：纵向变化量与横向变化量之比。"


def build_pdf(path: Path, text: str) -> None:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(f"BT /F1 14 Tf 72 700 Td ({text}) Tj ET".encode("latin-1"))).encode() + b" >>\nstream\n"
        + f"BT /F1 14 Tf 72 700 Td ({text}) Tj ET".encode("latin-1") + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    buffer = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(buffer))
        buffer += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(buffer)
    buffer += f"xref\n0 {len(objects) + 1}\n".encode()
    buffer += b"0000000000 65535 f \n"
    for offset in offsets:
        buffer += f"{offset:010d} 00000 n \n".encode()
    buffer += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    path.write_bytes(bytes(buffer))


def main() -> int:
    output_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp(prefix="omni-edu-parse-smoke-"))
    output_dir.mkdir(parents=True, exist_ok=True)

    engines = _handle({"schemaVersion": "xiazhi.parse.request.v1", "requestId": "engines-1", "type": "engines"})
    assert engines["ok"] is True, f"engines probe must succeed: {engines}"
    available = {entry["name"]: entry["available"] for entry in engines["engines"]}
    assert available.get("pypdf") is True, f"pypdf must be available: {engines}"

    pdf_path = output_dir / "parse-fixture.pdf"
    build_pdf(pdf_path, PDF_TEXT)
    pdf_result = _handle({"schemaVersion": "xiazhi.parse.request.v1", "requestId": "parse-pdf", "type": "parse", "filePath": str(pdf_path), "extension": ".pdf"})
    assert pdf_result["ok"] is True, f"pdf parse must succeed: {pdf_result}"
    assert PDF_TEXT.split(" y")[0] in pdf_result["text"], f"pdf text must be extracted: {pdf_result['text']!r}"
    assert pdf_result["engine"] in {"docling", "pypdf"}, f"pdf engine must be declared: {pdf_result['engine']}"
    assert pdf_result["pages"][0]["page"] == 1, "pdf pages must be numbered"

    docx_path = output_dir / "parse-fixture.docx"
    import docx

    document = docx.Document()
    document.add_paragraph(DOCX_TEXT)
    document.add_paragraph(DOCX_BODY)
    document.save(str(docx_path))
    docx_result = _handle({"schemaVersion": "xiazhi.parse.request.v1", "requestId": "parse-docx", "type": "parse", "filePath": str(docx_path), "extension": ".docx"})
    assert docx_result["ok"] is True, f"docx parse must succeed: {docx_result}"
    assert DOCX_TEXT in docx_result["text"] and DOCX_BODY in docx_result["text"], f"chinese docx text must survive parsing: {docx_result['text']!r}"

    missing = _handle({"schemaVersion": "xiazhi.parse.request.v1", "requestId": "parse-missing", "type": "parse", "filePath": str(output_dir / "nope.pdf"), "extension": ".pdf"})
    assert missing["ok"] is False and missing["code"] == "FILE_READ_FAILED", f"missing file must fail closed: {missing}"

    unsupported = _handle({"schemaVersion": "xiazhi.parse.request.v1", "requestId": "parse-xlsx", "type": "parse", "filePath": str(pdf_path), "extension": ".xlsx"})
    assert unsupported["ok"] is False and unsupported["code"] == "UNSUPPORTED_FORMAT", f"unsupported format must fail closed: {unsupported}"

    bad_schema = _handle({"schemaVersion": "wrong.v1", "requestId": "bad", "type": "engines"})
    assert bad_schema["ok"] is False and bad_schema["code"] == "BAD_SCHEMA", f"bad schema must fail closed: {bad_schema}"

    print(json.dumps({
        "ok": True,
        "pdfEngine": pdf_result["engine"],
        "pdfChars": len(pdf_result["text"]),
        "docxChars": len(docx_result["text"]),
        "doclingAvailable": bool(available.get("docling")),
        "errorPathsFailClosed": True,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
